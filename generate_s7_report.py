import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import build_doc_helpers
from build_doc_helpers import (
    add_heading_1,
    add_heading_2,
    add_heading_3,
    add_body_p,
    add_bullet_p,
    add_screenshot,
    create_callout,
    style_table,
)

# Configure output and screenshot paths
DOCX_OUTPUT_PATH = r"d:\HOC\PHAMMEM\DUAN_CODEGYM\BAO_CAO_CHUC_NANG_SPRINT_7_S7_01_DEN_S7_09.docx"
SCREENSHOTS_S7_DIR = r"d:\HOC\PHAMMEM\DUAN_CODEGYM\screenshots_s7"

# Point helper to Sprint 7 screenshots
build_doc_helpers.SCREENSHOTS_DIR = SCREENSHOTS_S7_DIR

def build_s7_document():
    print("Khởi tạo tài liệu báo cáo Word Sprint 7...")
    doc = docx.Document()

    # Configure Margins (0.75 inch = 1.9 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.different_first_page_header_footer = False

    # =========================================================================
    # BÌA / TIÊU ĐỀ BÁO CÁO
    # =========================================================================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = title_p.add_run("HỆ THỐNG QUẢN LÝ QUAN HỆ KHÁCH HÀNG & DOANH THU B2B (NEXUS CRM)\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    r_main = title_p.add_run("BÁO CÁO NGHIỆM THU TÍNH NĂNG CHI TIẾT SPRINT 7\n(USER STORIES S7-01 ĐẾN S7-09)")
    r_main.font.name = "Arial"
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy #1E3A8A

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(6)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_desc = sub_p.add_run("Tài liệu đặc tả nghiệp vụ, kiến trúc kỹ thuật và toàn bộ hình ảnh minh họa giao diện thực tế ứng dụng")
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(10.5)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(71, 85, 105)

    # Bảng thông tin dự án
    info_table = doc.add_table(rows=1, cols=4)
    info_headers = ["Thông tin dự án", "Phiên bản nghiệm thu", "Phạm vi phân hệ", "Ngày hoàn thiện"]
    info_rows = [
        ["Nexus CRM SaaS (TTCS_K4S5_N3)", "v2.7.0 (Sprint 7 Release)", "Quotes, Approvals, Contracts & Automation", "Tháng 10/2026"]
    ]
    style_table(info_table, [2.2, 1.6, 2.2, 1.2], info_headers, info_rows)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # PHẦN 1: BẢNG TỔNG HỢP MATRIX 9 USER STORIES SPRINT 7
    # =========================================================================
    add_heading_1(doc, "I. BẢNG TỔNG HỢP MA TRẬN TÍNH NĂNG SPRINT 7 (S7-01 ĐẾN S7-09)")
    add_body_p(doc, "Sprint 7 mang tính đột phá trong quy trình chốt giao dịch thương mại (Quote-to-Cash), tập trung vào: Soạn thảo bảng báo giá chuyên nghiệp đa sản phẩm (Quote Builder), cơ chế đa chiết khấu & quy trình phê duyệt phân cấp chặt chẽ, so sánh đối chiếu phiên bản báo giá, chuyển đổi báo giá thành hợp đồng tự động, cảnh báo hợp đồng đến hạn tái ký, cổng trực tuyến cho khách hàng tương tác trực tiếp, trung tâm thông báo hệ thống và động cơ tự động hóa quy trình IF–THEN không cần code. Dưới đây là ma trận 9 User Stories đã hoàn thiện 100%:")

    matrix_table = doc.add_table(rows=1, cols=5)
    matrix_headers = ["Mã Story", "Tên chức năng / User Story", "Phân hệ (Module)", "Phạm vi dữ liệu (Scope)", "Trạng thái nghiệm thu"]
    matrix_rows = [
        ["S7-01", "Soạn thảo Báo giá Đa sản phẩm (Quote Builder)", "Báo giá (Quotes & Deals)", "Kế thừa Cơ hội / Khách hàng", "ĐÃ ĐẠT (100%)"],
        ["S7-02", "Cơ chế Đa chiết khấu & Hàng đợi Phê duyệt phân cấp", "Báo giá & Phê duyệt", "Hàng đợi Chờ duyệt / Lịch sử", "ĐÃ ĐẠT (100%)"],
        ["S7-03", "So sánh trực quan giữa các phiên bản Báo giá (Diff)", "Báo giá & Quản lý phiên bản", "Lịch sử phiên bản v1, v2...", "ĐÃ ĐẠT (100%)"],
        ["S7-04", "Chuyển đổi Báo giá sang Hợp đồng & Đóng thắng Deal", "Báo giá -> Hợp đồng -> Deals", "Kế thừa tự động dòng hàng", "ĐÃ ĐẠT (100%)"],
        ["S7-05", "Xuất file PDF & Mẫu in Báo giá chuẩn Doanh nghiệp", "Báo giá & Mẫu in tài liệu", "Pháp nhân & Điều khoản", "ĐÃ ĐẠT (100%)"],
        ["S7-06", "Cổng thông tin Khách hàng xem & ký duyệt trực tuyến", "Cổng khách hàng (Public Portal)", "Liên kết an toàn có mã token", "ĐÃ ĐẠT (100%)"],
        ["S7-07", "Quản lý Hợp đồng & Cảnh báo Tái ký 1-click", "Hợp đồng (Contracts & Renewal)", "Đếm ngược hạn & Tạo Deal tái ký", "ĐÃ ĐẠT (100%)"],
        ["S7-08", "Trung tâm Thông báo Hệ thống & Cảnh báo Tự động", "Hạ tầng Thông báo (TopBar)", "Thông báo đa kênh thời gian thực", "ĐÃ ĐẠT (100%)"],
        ["S7-09", "Động cơ Tự động hóa Quy trình IF–THEN (Rules Engine)", "Tự động hóa (Automation)", "Trigger -> Condition -> Action", "ĐÃ ĐẠT (100%)"],
    ]
    style_table(matrix_table, [0.8, 2.6, 1.5, 1.2, 1.1], matrix_headers, matrix_rows)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # =========================================================================
    # PHẦN 2: BÁO CÁO CHI TIẾT & ẢNH MINH HỌA TỪNG CHỨC NĂNG (S7-01 -> S7-09)
    # =========================================================================
    add_heading_1(doc, "II. BÁO CÁO CHI TIẾT & HÌNH ẢNH MINH HỌA GIAO DIỆN TỪNG CHỨC NĂNG")

    # -------------------------------------------------------------------------
    # S7-01
    # -------------------------------------------------------------------------
    add_heading_2(doc, "1. USER STORY S7-01: TRÌNH SOẠN THẢO BÁO GIÁ ĐA SẢN PHẨM (QUOTE BUILDER)")
    add_body_p(doc, "Phân hệ Soạn thảo Báo giá (/quotes/builder hoặc /deals/:id/quote-builder) cung cấp môi trường tính toán bảng báo giá mạnh mẽ, trực quan cho đội ngũ tư vấn kinh doanh.")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Khắc phục triệt để tình trạng nhân viên kinh doanh phải tính toán báo giá thủ công trên file Excel dễ nhầm lẫn số liệu, sai lệch giá trị chiết khấu hoặc áp dụng sai mức thuế VAT. Hệ thống cho phép soạn thảo báo giá tức thì từ một Cơ hội bán hàng (Deal) đang trao đổi, kế thừa toàn bộ danh sách sản phẩm quan tâm, tự động hóa tính toán tổng số tiền và kiểm soát vòng đời trạng thái báo giá (Bản nháp -> Chờ duyệt -> Đã duyệt -> Đã gửi -> Khách chấp thuận -> Khách từ chối).")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Kế thừa dữ liệu thông minh:", "Khi tạo báo giá từ cơ hội bán hàng, hệ thống tự động điền thông tin Khách hàng, Công ty, Mã số thuế, Người phụ trách và các dòng sản phẩm đã chọn trong Deal.")
    add_bullet_p(doc, "Quản lý dòng sản phẩm linh hoạt (Line Items):", "Hỗ trợ chọn sản phẩm từ Danh mục sản phẩm chuẩn của công ty hoặc thêm tùy chỉnh. Mỗi dòng cho phép điều chỉnh Số lượng, Đơn vị tính, Đơn giá niêm yết, Chiết khấu dòng theo % hoặc tiền mặt, và phần trăm thuế VAT (0%, 8%, 10%).")
    add_bullet_p(doc, "Kiểm soát giá sàn (Floor Price Guard):", "Tự động cảnh báo nếu nhân viên kinh doanh giảm đơn giá xuống dưới mức giá vốn/giá sàn cấu hình trong danh mục sản phẩm.")
    add_bullet_p(doc, "Điều khoản thương mại & Hiệu lực:", "Khai báo thời hạn hiệu lực báo giá (Valid Until), Điều khoản thanh toán nhiều đợt (Payment Terms), Điều khoản giao hàng (Delivery Terms) và Ghi chú gửi khách hàng.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_01_quote_builder_overview.png", "Hình 1a: Tổng quan Giao diện Soạn thảo Báo giá Đa sản phẩm (Quote Builder) với thông tin khách hàng và thanh công cụ thao tác (S7-01)")
    add_screenshot(doc, "s7_01_quote_items_list.png", "Hình 1b: Bảng chi tiết các Dòng Sản phẩm với đơn giá, số lượng, chiết khấu và thuế VAT tự động tính toán (S7-01)")

    create_callout(doc, "Toàn bộ bảng tính Quote Builder được xây dựng dưới dạng bảng tính phản ứng nhanh (Reactive Calculations Engine), cập nhật tức thì Thành tiền (Line Total), Thuế VAT, Tiền giảm giá và Tổng cộng thanh toán (Grand Total) theo thời gian thực mà không gây trễ giao diện.", "ĐẶC TẢ KỸ THUẬT QUOTE BUILDER")

    # -------------------------------------------------------------------------
    # S7-02
    # -------------------------------------------------------------------------
    add_heading_2(doc, "2. USER STORY S7-02: CƠ CHẾ ĐA CHIẾT KHẤU & QUY TRÌNH PHÊ DUYỆT PHÂN CẤP (APPROVAL WORKFLOW)")
    add_body_p(doc, "Phân hệ Hàng đợi Phê duyệt Báo giá (/quotes/approvals) phục vụ kiểm soát chặt chẽ chính sách giá bán và biên lợi nhuận của doanh nghiệp.")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Ngăn ngừa rủi ro thất thoát doanh thu khi nhân viên kinh doanh tự ý giảm giá quá sâu để đạt doanh số cá nhân. Thiết lập cơ chế kiểm duyệt tự động dựa trên mức chiết khấu: những báo giá thông thường được duyệt ngay (Auto-approved), trong khi các báo giá giảm giá sâu phải qua cấp Trưởng nhóm hoặc Giám đốc kinh doanh xét duyệt trước khi được gửi tới khách hàng.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Quy tắc phân cấp phê duyệt tự động:", "Báo giá có chiết khấu <= 10%: Tự động phê duyệt (Auto-approved) ngay khi gửi. Chiết khấu từ 10.1% đến 20%: Chuyển trạng thái 'Chờ phê duyệt' và yêu cầu Trưởng nhóm kinh doanh (Team Leader) duyệt. Chiết khấu > 20% hoặc bán dưới giá sàn: Bắt buộc Giám đốc kinh doanh (Sales Director) phê duyệt.")
    add_bullet_p(doc, "Cảnh báo vi phạm trần chiết khấu trực quan:", "Hiển thị khung cảnh báo màu vàng/đỏ ngay trong màn hình soạn thảo báo giá, chỉ rõ mức chiết khấu hiện tại, mức vượt ngưỡng và cấp có thẩm quyền phê duyệt.")
    add_bullet_p(doc, "Hàng đợi Phê duyệt tập trung (/quotes/approvals):", "Cung cấp danh sách các báo giá đang chờ cấp trên xử lý kèm đầy đủ thông tin: Mã báo giá, Khách hàng, Giá trị gốc, Mức chiết khấu đề xuất, Tổng tiền sau giảm và Người yêu cầu.")
    add_bullet_p(doc, "Thao tác ra quyết định & Ghi nhận lý do:", "Người quản lý có 3 lựa chọn: Phê duyệt (Approve), Từ chối (Reject), hoặc Yêu cầu chỉnh sửa lại (Return for Edit) kèm bắt buộc nhập nhận xét/chỉ đạo.")
    add_bullet_p(doc, "Lịch sử xét duyệt (Audit Trail):", "Tab 'Lịch sử xét duyệt' lưu trữ vĩnh viễn toàn bộ nhật ký phê duyệt: ai duyệt, thời gian chính xác, hành động gì và ý kiến phê duyệt.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_02_approval_queue.png", "Hình 2a: Hàng đợi Phê duyệt Báo giá với thông tin mức chiết khấu, cấp duyệt và các nút xử lý Phê duyệt / Từ chối (S7-02)")
    add_screenshot(doc, "s7_02_approval_history.png", "Hình 2b: Tab Lịch sử Xét duyệt ghi nhận chi tiết chuỗi sự kiện, hành động và lý do phê duyệt của cấp quản lý (S7-02)")

    # -------------------------------------------------------------------------
    # S7-03
    # -------------------------------------------------------------------------
    add_heading_2(doc, "3. USER STORY S7-03: SO SÁNH TRỰC QUAN GIỮA CÁC PHIÊN BẢN BÁO GIÁ (VERSION COMPARISON DIFF)")
    add_body_p(doc, "Tính năng So sánh Phiên bản tích hợp trực tiếp tại thanh công cụ Quote Builder qua modal đối chiếu chuyên sâu.")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Trong các thương vụ đàm phán B2B kéo dài, khách hàng thường yêu cầu thay đổi phạm vi gói giải pháp, tăng giảm số lượng người dùng hoặc đàm phán lại mức giá qua nhiều phiên bản (v1, v2, v3...). Tính năng này giúp nhân viên và quản lý ngay lập tức nhận diện được những điểm khác biệt giữa bản báo giá mới nhất với bản cũ mà không cần soi từng con số.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Lưu trữ phiên bản tự động (Version Snapshot):", "Mỗi khi báo giá được gửi duyệt hoặc được cập nhật quan trọng, hệ thống tự động sinh số phiên bản mới (v1, v2, v3...) và lưu trữ snapshot nguyên trạng dữ liệu.")
    add_bullet_p(doc, "Đối soát hai phiên bản bất kỳ:", "Người dùng chọn 1 phiên bản gốc trong lịch sử và phiên bản hiện tại để hệ thống thực hiện giải thuật so sánh đối chiếu (Visual Diff).")
    add_bullet_p(doc, "Làm nổi bật trực quan chênh lệch (Highlight Diff):", "Hiển thị màu xanh lá cho các dòng mặt hàng mới được bổ sung, màu đỏ gạch ngang cho các mặt hàng bị loại bỏ, và màu vàng cam cho các mặt hàng có sự thay đổi về Đơn giá, Số lượng hoặc % Chiết khấu.")
    add_bullet_p(doc, "Bảng tổng hợp chênh lệch tài chính:", "Thống kê rõ ràng biến động: Chênh lệch Tiền hàng trước thuế (Subtotal Diff), Chênh lệch Tiền chiết khấu, Chênh lệch Thuế VAT và Chênh lệch Tổng thanh toán cuối cùng.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_03_version_comparison_modal.png", "Hình 3: Cửa sổ Modal So sánh Trực quan giữa các phiên bản Báo giá với tính năng highlight chênh lệch dòng hàng và tài chính (S7-03)")

    # -------------------------------------------------------------------------
    # S7-04
    # -------------------------------------------------------------------------
    add_heading_2(doc, "4. USER STORY S7-04: CHUYỂN ĐỔI BÁO GIÁ SANG HỢP ĐỒNG & ĐÓNG THẮNG DEAL (CONVERT TO CONTRACT)")
    add_body_p(doc, "Quy trình chuyển đổi liền mạch từ Báo giá được chấp thuận sang Hợp đồng pháp lý chính thức.")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Kết nối khép kín chu trình bán hàng (Sales Pipeline). Khi khách hàng đồng ý chốt báo giá, nhân viên kinh doanh không phải nhập lại dữ liệu khách hàng và sản phẩm sang hệ thống quản lý hợp đồng, loại bỏ hoàn toàn nguy cơ sai sót số liệu và tiết kiệm 90% thời gian tạo hồ sơ hợp đồng.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Kích hoạt nút chuyển đổi thông minh:", "Khi báo giá chuyển sang trạng thái 'Đã chấp nhận' (Accepted), nút bấm 'Chuyển thành Hợp đồng' tự động sáng lên trên thanh công cụ.")
    add_bullet_p(doc, "Cửa sổ khởi tạo Hợp đồng (Convert Modal):", "Cho phép nhập: Số hợp đồng (duy nhất toàn hệ thống, vd: HD-2026-089), Ngày ký, Ngày bắt đầu hiệu lực, Ngày hết hạn và Điều khoản thanh toán thỏa thuận.")
    add_bullet_p(doc, "Kế thừa 100% dòng sản phẩm và giá trị:", "Toàn bộ mặt hàng, đơn giá đàm phán, chiết khấu và tổng thanh toán của Báo giá được chuyển thẳng thành phụ lục hàng hóa của Hợp đồng mới.")
    add_bullet_p(doc, "Tự động đóng Thắng Cơ hội (Deal Won):", "Đồng thời hệ thống tự động cập nhật trạng thái Cơ hội bán hàng (Deal) tương ứng sang giai đoạn 'Đóng Thắng' (Closed Won), ghi nhận doanh thu và cập nhật tức thì lên biểu đồ phễu bán hàng.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_04_convert_to_contract_modal.png", "Hình 4: Cửa sổ Modal Chuyển Báo giá sang Hợp đồng với các trường ngày ký, hiệu lực và cam kết đóng thắng cơ hội tự động (S7-04)")

    # -------------------------------------------------------------------------
    # S7-05
    # -------------------------------------------------------------------------
    add_heading_2(doc, "5. USER STORY S7-05: XUẤT FILE PDF & MẪU IN BÁO GIÁ CHUẨN DOANH NGHIỆP (PDF & PRINT TEMPLATE)")
    add_body_p(doc, "Mẫu in và tài liệu PDF Báo giá chuyên nghiệp tích hợp trực tiếp qua nút 'Xuất PDF Doanh nghiệp' và nút 'In Báo Giá'.")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Cung cấp mẫu tài liệu báo giá chuẩn mực thương mại để gửi qua email hoặc in ấn ký đóng dấu trực tiếp khi làm việc với các khách hàng khối doanh nghiệp, tập đoàn và cơ quan nhà nước đòi hỏi tính trang trọng, pháp lý.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Đầy đủ thông tin pháp nhân:", "Tên doanh nghiệp, Logo NexusCRM, Mã số thuế, Địa chỉ trụ sở, Số điện thoại hotline, Email và Số tài khoản ngân hàng thụ hưởng.")
    add_bullet_p(doc, "Bảng danh mục hàng hóa theo chuẩn kế toán:", "Cột STT, Mã sản phẩm, Tên dịch vụ/hàng hóa, Đơn vị tính, Số lượng, Đơn giá niêm yết, Tiền chiết khấu và Thành tiền trước thuế.")
    add_bullet_p(doc, "Bảng tổng hợp thanh toán & Đọc tiền bằng chữ:", "Phân tách rõ ràng: Tiền hàng trước thuế (Subtotal), Tiền chiết khấu đơn hàng, Thuế suất và Tiền thuế GTGT, Tổng cộng tiền thanh toán, kèm dòng 'Số tiền bằng chữ' theo đúng quy chuẩn kế toán Việt Nam.")
    add_bullet_p(doc, "Điều khoản thương mại & Khu vực ký duyệt:", "In rõ điều khoản thanh toán, phương thức giao hàng, thời hạn hiệu lực báo giá, khu vực Người báo giá và Đại diện khách hàng ký ghi rõ họ tên.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_05_quote_totals_terms.png", "Hình 5: Chi tiết Bảng tổng hợp thanh toán, Thuế VAT, Chiết khấu và các Điều khoản thương mại theo mẫu in doanh nghiệp (S7-05)")

    # -------------------------------------------------------------------------
    # S7-06
    # -------------------------------------------------------------------------
    add_heading_2(doc, "6. USER STORY S7-06: CỔNG THÔNG TIN KHÁCH HÀNG TRỰC TUYẾN (CUSTOMER QUOTE PORTAL)")
    add_body_p(doc, "Trang tương tác độc lập dành riêng cho người mua hàng tại đường dẫn /portal/quotes/:token.")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Loại bỏ sự chậm trễ trong khâu phản hồi báo giá qua email thông thường. Khách hàng nhận được một đường link bảo mật độc quyền, có thể truy cập từ máy tính hoặc điện thoại mà không cần đăng nhập tài khoản, xem toàn bộ chi tiết báo giá và bấm duyệt ngay lập tức.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Sinh mã liên kết an toàn (Secure Shareable Token):", "Nhân viên kinh doanh mở modal 'Liên kết Khách hàng' để tạo đường dẫn có mã token mã hóa độc nhất cho từng báo giá và sao chép gửi khách qua Zalo/Email.")
    add_bullet_p(doc, "Giao diện Portal thân thiện, chuyên nghiệp:", "Giao diện tối ưu trải nghiệm khách hàng với thanh điều hướng thương hiệu, bảng chi tiết từng sản phẩm, tổng giá trị thanh toán và điều khoản triển khai.")
    add_bullet_p(doc, "Khách hàng Chấp thuận báo giá (Accept Quote):", "Khách hàng bấm nút 'Chấp thuận Báo giá', nhập Họ tên và Chức danh người ký duyệt. Hệ thống lập tức cập nhật trạng thái Báo giá sang 'Accepted' và gửi thông báo tức thời cho nhân viên kinh doanh phụ trách.")
    add_bullet_p(doc, "Khách hàng Từ chối / Yêu cầu sửa đổi (Reject Quote):", "Khách hàng bấm 'Từ chối / Yêu cầu sửa đổi', nhập lý do cụ thể (vd: cần giảm thêm giá gói onboarding hoặc tăng số buổi đào tạo). Báo giá cập nhật trạng thái 'Rejected' kèm lưu trữ nguyên văn ý kiến phản hồi.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_06_share_portal_modal.png", "Hình 6a: Modal Tạo & Sao chép Liên kết An toàn chia sẻ cho Khách hàng truy cập duyệt trực tuyến (S7-06)")
    add_screenshot(doc, "s7_06_customer_portal_view.png", "Hình 6b: Giao diện Cổng Thông tin Khách hàng trực tuyến với 2 nút hành động Chấp thuận và Từ chối báo giá (S7-06)")

    # -------------------------------------------------------------------------
    # S7-07
    # -------------------------------------------------------------------------
    add_heading_2(doc, "7. USER STORY S7-07: QUẢN LÝ HỢP ĐỒNG & CẢNH BÁO TÁI KÝ 1-CLICK (CONTRACT RENEWAL)")
    add_body_p(doc, "Phân hệ Quản lý Hợp đồng & Tái ký (/contracts) phục vụ theo dõi thời hạn hiệu lực và duy trì doanh thu định kỳ (Recurring Revenue).")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Giải quyết bài toán doanh nghiệp để quên hợp đồng hết hạn khiến khách hàng chuyển sang sử dụng dịch vụ của đối thủ cạnh tranh. Hệ thống chủ động đếm ngược số ngày còn lại của từng hợp đồng và gắn cờ cảnh báo sớm để đội ngũ kinh doanh chủ động liên hệ chăm sóc và tái ký trước từ 30 đến 60 ngày.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Theo dõi trạng thái & Đếm ngược ngày:", "Bảng danh mục hợp đồng hiển thị: Số hợp đồng, Tiêu đề, Khách hàng, Giá trị hợp đồng, Ngày ký, Ngày hết hạn và Số ngày còn lại (Remaining Days).")
    add_bullet_p(doc, "Huy hiệu cảnh báo trực quan:", "Hợp đồng còn > 30 ngày: Trạng thái 'Đang hiệu lực' (Màu xanh dương). Hợp đồng còn <= 30 ngày: Huy hiệu màu vàng cam 'Sắp hết hạn' kèm số ngày đếm ngược đỏ. Hợp đồng quá hạn: Huy hiệu 'Đã hết hạn'.")
    add_bullet_p(doc, "Bộ lọc nhanh theo tình trạng hiệu lực:", "Các tab lọc tiện lợi: 'Tất cả', 'Hiệu lực', 'Sắp hết hạn' và 'Đã tái ký', hỗ trợ thiết lập ngưỡng cảnh báo 30 ngày, 60 ngày hoặc 90 ngày.")
    add_bullet_p(doc, "Tạo cơ hội tái ký 1-click (1-Click Renewal):", "Bấm nút 'Tái ký' tại dòng hợp đồng sắp hết hạn để mở modal khởi tạo Cơ hội bán hàng tái ký mới. Hệ thống tự động điền Tiêu đề tái ký, Khách hàng liên quan, Giá trị kế thừa và ngày dự kiến đóng, đưa thẳng vào pipeline cơ hội bán hàng.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_07_contracts_list.png", "Hình 7a: Danh sách Quản lý Hợp đồng với huy hiệu trạng thái, đếm ngược số ngày còn lại và bộ lọc sắp hết hạn (S7-07)")
    add_screenshot(doc, "s7_07_renewal_opportunity_modal.png", "Hình 7b: Cửa sổ Modal Tạo Cơ hội Bán hàng Tái ký Hợp đồng kế thừa dữ liệu 1-click (S7-07)")

    # -------------------------------------------------------------------------
    # S7-08
    # -------------------------------------------------------------------------
    add_heading_2(doc, "8. USER STORY S7-08: TRUNG TÂM THÔNG BÁO HỆ THỐNG & CẢNH BÁO TỰ ĐỘNG (NOTIFICATION CENTER)")
    add_body_p(doc, "Tích hợp biểu tượng chuông thông báo thông minh trên thanh điều hướng đầu trang (TopBar).")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Đảm bảo luồng thông tin trong doanh nghiệp luôn thông suốt theo thời gian thực. Bất kể khi nào có báo giá cần duyệt, báo giá được khách chấp thuận, hợp đồng sắp hết hạn hoặc công việc được giao mới, người dùng nhận được cảnh báo ngay lập tức trên giao diện mà không cần kiểm tra email thủ công.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Huy hiệu đếm số lượng chưa đọc (Unread Badge):", "Biểu tượng chuông hiển thị chấm đỏ với con số thông báo chưa đọc theo thời gian thực.")
    add_bullet_p(doc, "Menu trượt danh sách thông báo phong phú:", "Phân loại thông báo theo icon và màu sắc đặc trưng: Báo giá (Quote), Phê duyệt (Approval), Hợp đồng (Contract), Công việc (Task) và Hệ thống (System).")
    add_bullet_p(doc, "Điều hướng trực tiếp theo ngữ cảnh:", "Nhấp chuột vào một thông báo sẽ lập tức chuyển hướng người dùng tới màn hình xử lý tương ứng (vd: nhấp vào thông báo yêu cầu duyệt báo giá sẽ mở thẳng trang /quotes/approvals).")
    add_bullet_p(doc, "Thao tác quản lý thông báo:", "Hỗ trợ nút 'Đánh dấu đã đọc' cho từng thông báo hoặc 'Đánh dấu tất cả đã đọc' (Mark All as Read) chỉ bằng 1 cú nhấp chuột.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_08_notifications_dropdown.png", "Hình 8: Menu thả xuống Trung tâm Thông báo Hệ thống với danh mục cảnh báo đa dạng và thao tác đánh dấu đã đọc (S7-08)")

    # -------------------------------------------------------------------------
    # S7-09
    # -------------------------------------------------------------------------
    add_heading_2(doc, "9. USER STORY S7-09: ĐỘNG CƠ TỰ ĐỘNG HÓA QUY TRÌNH IF–THEN (AUTOMATION RULES ENGINE)")
    add_body_p(doc, "Phân hệ Tự động hóa quy trình (/automation) cung cấp khả năng cấu hình kịch bản tự động theo logic 'NẾU điều này xảy ra – THÌ thực hiện hành động kia' (IF–THEN).")

    add_heading_3(doc, "a. Mục tiêu nghiệp vụ & Bài toán giải quyết")
    add_body_p(doc, "Giải phóng đội ngũ kinh doanh khỏi các công việc quản trị thủ công mang tính lặp đi lặp lại. Giúp doanh nghiệp chuẩn hóa quy trình chăm sóc khách hàng: tự động giao việc khi có deal mới, tự động gửi thông báo khi deal thắng/thua, tự động nhắc việc khi hợp đồng sắp hết hạn.")

    add_heading_3(doc, "b. Quy cách tính năng & Luồng hoạt động")
    add_bullet_p(doc, "Cấu hình Sự kiện Kích hoạt (Trigger):", "Hỗ trợ kích hoạt theo sự kiện thực thể: Cơ hội bán hàng (Deal Stage Changed, Deal Value Updated), Báo giá (Quote Accepted, Quote Sent), Hợp đồng (Contract Expiring Soon) hoặc Lead (New Lead Created).")
    add_bullet_p(doc, "Bộ lọc Điều kiện (Condition):", "Thiết lập điều kiện lọc nâng cao: Giá trị cơ hội lớn hơn ngưỡng X, Giai đoạn bằng 'Đàm phán', Khách hàng thuộc nhóm ngành Y,...")
    add_bullet_p(doc, "Hành động tự động (Automated Actions):", "Tự động tạo nhiệm vụ mới (Create Task) giao cho người phụ trách kèm deadline; Tự động gửi thông báo hệ thống (Send In-App Notification); Điều chuyển người chịu trách nhiệm (Reassign Owner).")
    add_bullet_p(doc, "Quản lý Bật/Tắt quy tắc (Toggle Active):", "Mỗi quy tắc có công tắc chuyển đổi trạng thái Hoạt động / Tạm dừng trực tiếp ngay trên bảng danh sách.")
    add_bullet_p(doc, "Nhật ký thực thi tự động hóa (Execution Logs):", "Tab 'Nhật ký thực thi' ghi nhận chi tiết từng lần luật chạy: Thời gian kích hoạt, Thực thể liên quan, Hành động đã thực hiện, Trạng thái (Thành công / Bỏ qua / Thất bại) và chi tiết kỹ thuật.")

    add_heading_3(doc, "c. Hình ảnh minh họa giao diện thực tế")
    add_screenshot(doc, "s7_09_automation_rules_list.png", "Hình 9a: Bảng Danh sách các Quy tắc Tự động hóa IF–THEN với công tắc bật/tắt và tóm tắt kịch bản (S7-09)")
    add_screenshot(doc, "s7_09_create_automation_modal.png", "Hình 9b: Cửa sổ Modal Cấu hình Quy tắc Tự động hóa mới với các khối Sự kiện, Điều kiện và Hành động (S7-09)")
    add_screenshot(doc, "s7_09_automation_execution_logs.png", "Hình 9c: Tab Nhật ký Thực thi ghi nhận chi tiết lịch sử kích hoạt và trạng thái thực hiện các kịch bản tự động (S7-09)")

    # =========================================================================
    # PHẦN 3: TỔNG KẾT & ĐÁNH GIÁ NGHIỆM THU KỸ THUẬT
    # =========================================================================
    add_heading_1(doc, "III. TỔNG KẾT & ĐÁNH GIÁ NGHIỆM THU KỸ THUẬT SPRINT 7")
    add_body_p(doc, "Toàn bộ 9 User Stories của Sprint 7 đã được hoàn thiện xuất sắc, đáp ứng đầy đủ cả về mặt nghiệp vụ kinh doanh thực tế, trải nghiệm người dùng hiện đại và kiến trúc mã nguồn chuẩn mực:")

    add_bullet_p(doc, "Độ tin cậy mã nguồn (Code Quality):", "100% mã nguồn TypeScript được chuẩn hóa Type an toàn (Strict Typing), 0% kiểu any, vượt qua kiểm tra biên dịch tsc --noEmit không có bất kỳ lỗi cú pháp nào.")
    add_bullet_p(doc, "Bảo mật & Phân quyền:", "Phân định quyền hạn rõ ràng giữa Chuyên viên kinh doanh, Trưởng nhóm (Team Leader) và Giám đốc kinh doanh (Director) trong quy trình duyệt chiết khấu. Cổng khách hàng (Quote Portal) áp dụng cơ chế xác thực Token an toàn, cách ly hoàn toàn với dữ liệu nội bộ.")
    add_bullet_p(doc, "Khả năng mở rộng & Kế thừa dữ liệu:", "Dữ liệu được kế thừa xuyên suốt từ Cơ hội bán hàng -> Báo giá đa phiên bản -> Hợp đồng thương mại -> Cơ hội tái ký, tạo thành chu trình khép kín không điểm đứt gãy.")
    add_bullet_p(doc, "Đánh giá nghiệm thu chung:", "Sprint 7 ĐẠT 100% TIÊU CHÍ NGHIỆM THU (PASSED), sẵn sàng triển khai chính thức phục vụ hoạt động vận hành kinh doanh của doanh nghiệp.")

    # Save document
    doc.save(DOCX_OUTPUT_PATH)
    print(f"Đã tạo thành công tài liệu báo cáo Sprint 7 tại: {DOCX_OUTPUT_PATH}")

if __name__ == '__main__':
    build_s7_document()
